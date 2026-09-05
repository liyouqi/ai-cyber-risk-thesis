"""Extract the AI Act technical checklist into the common experiment format."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("data/AI_Tool_Onboarding_Risk_Assessment.xlsx")
OUTPUT = Path("experiments/datasets/ai_act/items.csv")
SHEET = "技术控制清单IT Control CheckList"
HEADER_ROW = 4
KNOWN_TRANSLATIONS = {
    "生成式AI/RAG/Agent": "Generative AI, RAG or agent systems",
}


def english(value: object) -> str:
    text = "" if value is None else str(value).strip()
    if text in KNOWN_TRANSLATIONS:
        return KNOWN_TRANSLATIONS[text]
    _, separator, translation = text.partition(" / ")
    return translation.strip() if separator else text


def main() -> None:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET]
        rows = []
        for row_number in range(HEADER_ROW + 1, sheet.max_row + 1):
            item_id = english(sheet.cell(row_number, 1).value)
            if not item_id.startswith("TC"):
                continue
            rows.append(
                {
                    "item_id": item_id,
                    "framework": "EU AI Act",
                    "domain": english(sheet.cell(row_number, 2).value),
                    "control_statement": english(sheet.cell(row_number, 3).value),
                    "expected_evidence": english(sheet.cell(row_number, 4).value),
                    "applicability": english(sheet.cell(row_number, 7).value),
                    "existing_mapping": english(sheet.cell(row_number, 8).value),
                    "source_row": row_number,
                }
            )
    finally:
        workbook.close()

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    print(f"Wrote {len(rows)} items to {OUTPUT}")


if __name__ == "__main__":
    main()
